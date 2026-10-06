#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
量化基金生存录 · 玩家数据分析

用法（二选一）：
  python3 analyze.py --url https://xiaolipearl.com --token 你的ADMIN_TOKEN
  python3 analyze.py --file raw.json            # 用之前下载好的原始数据

会生成：
  raw_YYYYMMDD.json        原始记录（每局一行）
  玩家数据报告_YYYYMMDD.xlsx  分析报告（概览、漏斗、每年选择、人格、军师、成绩、思考时长、下期素材、原始数据）
需要：pip install openpyxl
"""
import argparse, json, sys, time, statistics as st, urllib.request, urllib.parse
from collections import Counter, defaultdict
from datetime import datetime

META = {"opts": [["防御", "顺势而为", "中性"], ["追价值", "开辟新阿尔法", "押小盘"], ["研究财务质量因子", "跟随趋势", "抄底"], ["保持防御", "中性", "押拐点"], ["扩张", "降成本", "差异化"], ["建拥挤度监控系统", "放大仓位", "建机器学习平台"], ["适度扩张", "发新产品", "封盘"], ["维持现状", "主动去拥挤", "整体降仓位"], ["风控优先", "顺势", "抄底"], ["押拐点", "保持动量", "中性"]], "reals": ["互联网泡沫破裂", "经济衰退 · 9·11", "安然、世通造假", "熊市见底大反弹", "市场温和复苏", "加息 · 横盘", "房价见顶", "8 月「量化地震」", "雷曼倒闭 · 金融危机", "零利率 · 动量崩溃"], "adv": [["小林", "阿杰", "老周"], ["阿杰", "老周", "小林"], ["小林", "阿杰", "Lily"], ["阿杰", "老周", "小林"], ["Lily", "老周", "小林"], ["老周", "阿杰", "小林"], ["小林", "Lily", "老周"], ["阿杰", "老周", "小林"], ["老周", "阿杰", "王总"], ["小林", "阿杰", "老周"]], "persona": {"TA": "风口冲浪者", "TD": "稳健跟随者", "CA": "抄底狙击手", "CD": "冷静的反向派"}, "ch": {"jie": "阿杰", "zhou": "老周", "lin": "小林", "lily": "Lily", "wang": "王总", "aunt": "张阿姨"}, "best": [0, 2, 0, 2, 2, 0, 2, 1, 0, 0], "worst": [1, 1, 0, 0, 0, 0, 0, 0, 0, 0]}
GY0 = 2027
CODE_DIM = {"T": "顺势", "C": "逆向", "A": "进攻", "D": "防守", "L": "长期", "Q": "灵活", "O": "专一", "M": "兼听"}
SRC_NAME = {"dy": "抖音", "bl": "B站", "xhs": "小红书", "wx": "微信", "": "直接打开"}

def fetch(url, token):
    items, cursor = [], ""
    while True:
        q = {"token": token}
        if cursor: q["cursor"] = cursor
        u = url.rstrip("/") + "/api/export?" + urllib.parse.urlencode(q)
        d = json.load(urllib.request.urlopen(u, timeout=60))
        if not d.get("ok"): sys.exit("导出失败：口令不对，或 KV 没有绑定")
        items += d["items"]; cursor = d.get("cursor", "")
        print(f"  已下载 {len(items)} 条", end="\r")
        if d.get("complete") or not cursor: break
    print()
    return items

def pct(a, b): return a / b if b else 0

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--url"); ap.add_argument("--token"); ap.add_argument("--file")
    ap.add_argument("--since", help="只分析这一天之后的数据，如 20261010")
    a = ap.parse_args()
    if a.file: rows = json.load(open(a.file, encoding="utf-8"))
    elif a.url and a.token:
        rows = fetch(a.url, a.token)
        fn = f"raw_{datetime.now():%Y%m%d}.json"; json.dump(rows, open(fn, "w", encoding="utf-8"), ensure_ascii=False); print("原始数据已保存：", fn)
    else: ap.error("需要 --url 和 --token，或 --file")
    rows = [r for r in rows if r.get("id") and not str(r.get("id")).startswith("seed")]
    if a.since: rows = [r for r in rows if str(r.get("day", "")) >= a.since]
    if not rows: sys.exit("没有数据")
    build(rows)

def build(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, Reference
    F = "Arial"; H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="3E3833")
    B = Font(name=F, bold=True); N = Font(name=F); T = Font(name=F, bold=True, size=14); G = Font(name=F, color="808080", italic=True)
    wb = Workbook()
    def sheet(name, title, note=None, first=False):
        ws = wb.active if first else wb.create_sheet(name); ws.title = name
        ws["A1"] = title; ws["A1"].font = T
        if note: ws["A2"] = note; ws["A2"].font = G
        return ws
    def head(ws, r, cols):
        for i, c in enumerate(cols, 1):
            x = ws.cell(row=r, column=i, value=c); x.font = H; x.fill = HF; x.alignment = Alignment(horizontal="center")
    def put(ws, r, vals, fmts=None):
        for i, v in enumerate(vals, 1):
            x = ws.cell(row=r, column=i, value=v); x.font = N
            if fmts and i - 1 < len(fmts) and fmts[i - 1]: x.number_format = fmts[i - 1]
    def widths(ws, ws_):
        for i, w in enumerate(ws_, 1): ws.column_dimensions[chr(64 + i)].width = w

    n_all = len(rows)
    ends = [r for r in rows if r.get("ev") == "end"]
    full = [r for r in ends if not r.get("dead")]
    dead = [r for r in ends if r.get("dead")]
    reach = Counter(len(r.get("ch", "")) for r in rows)
    durs = [r["dur"] for r in ends if r.get("dur")]
    devs = Counter(r.get("dev", "") for r in rows)
    srcs = Counter(r.get("src", "") for r in rows)
    replay = sum(1 for r in rows if (r.get("run") or 1) > 1)

    # 概览
    ws = sheet("概览", "量化基金生存录 · 玩家数据概览", f"数据范围：{min(r['day'] for r in rows)} – {max(r['day'] for r in rows)}；一局 = 一条记录（同一个人玩多次会算多局）", first=True)
    kv = [("开始游戏的局数", n_all, "#,##0"), ("玩到结局的局数", len(ends), "#,##0"), ("完成率", pct(len(ends), n_all), "0.0%"),
          ("中途被清盘的比例（占完成）", pct(len(dead), len(ends)), "0.0%"), ("平均游戏时长（分钟）", (st.mean(durs) / 60) if durs else 0, "0.0"),
          ("中位游戏时长（分钟）", (st.median(durs) / 60) if durs else 0, "0.0"), ("重玩的局数占比", pct(replay, n_all), "0.0%"),
          ("手机端占比", pct(devs.get("m", 0), n_all), "0.0%"),
          ("平均年化超额（玩满十年）", st.mean([r["ann"] for r in full]) if full else 0, "0.00%"),
          ("张阿姨 10 万的平均结局（万）", st.mean([r["aunt"] for r in ends if r.get("aunt") is not None]) if ends else 0, "0.0"),
          ("跑赢指数的比例（10 万 > 买纳指的 5.6 万）", pct(sum(1 for r in ends if (r.get("aunt") or 0) > 5.58), len(ends)), "0.0%"),
          ("守住本金的比例（10 万 ≥ 10 万）", pct(sum(1 for r in ends if (r.get("aunt") or 0) >= 10), len(ends)), "0.0%")]
    head(ws, 4, ["指标", "数值"])
    for i, (k, v, f) in enumerate(kv): put(ws, 5 + i, [k, v], [None, f])
    r0 = 6 + len(kv); head(ws, r0, ["来源（链接里的 ?from=）", "局数", "占比"])
    for i, (k, v) in enumerate(srcs.most_common()): put(ws, r0 + 1 + i, [SRC_NAME.get(k, k), v, pct(v, n_all)], [None, "#,##0", "0.0%"])
    widths(ws, [40, 14, 10])

    # 漏斗
    ws = sheet("流失漏斗", "每一年还有多少人在玩", "「到达第 N 年」= 至少做完了 N 个决定")
    head(ws, 4, ["年份", "真实对应", "做完这一年的局数", "占开始的比例"])
    for y in range(10):
        c = sum(v for k, v in reach.items() if k >= y + 1)
        put(ws, 5 + y, [f"{GY0 + y}", META["reals"][y], c, pct(c, n_all)], [None, None, "#,##0", "0.0%"])
    ch = BarChart(); ch.title = "流失漏斗"; ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=14), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=14)); ch.legend = None; ch.height = 7; ch.width = 16; ws.add_chart(ch, "F4")
    widths(ws, [10, 22, 18, 14])

    # 每年选择
    ws = sheet("每年选择", "每一年，大家选了什么", "最优路径 = 59,049 种玩法里第一名的那条路；★ 标出这一年最优路径的选择")
    head(ws, 4, ["游戏年份", "真实事件", "选项", "主张者", "选的局数", "占比", "最优路径"])
    r = 5; yearpick = []
    for y in range(10):
        cnt = Counter(x["ch"][y] for x in rows if len(x.get("ch", "")) > y); tot = sum(cnt.values())
        for j, L in enumerate("ABC"):
            star = "★" if META["best"][y] == j else ""
            put(ws, r, [f"{GY0 + y}" if j == 0 else "", META["reals"][y] if j == 0 else "", f"{L}. {META['opts'][y][j]}", META["adv"][y][j], cnt.get(L, 0), pct(cnt.get(L, 0), tot), star], [None, None, None, None, "#,##0", "0.0%"])
            r += 1
        yearpick.append((y, cnt, tot)); r += 1
    widths(ws, [10, 22, 22, 10, 10, 10, 10])

    # 人格
    ws = sheet("投资人格", "投资人格分布（玩到结局的局）")
    codes = Counter(x.get("code", "") for x in ends if x.get("code")); base = Counter(c[:2] for c in codes.elements())
    head(ws, 4, ["人格", "名称", "局数", "占比"])
    for i, (k, v) in enumerate(base.most_common()): put(ws, 5 + i, [k, META["persona"].get(k, ""), v, pct(v, len(ends))], [None, None, "#,##0", "0.0%"])
    r0 = 6 + len(base); head(ws, r0, ["完整代码", "含义", "局数", "占比"])
    for i, (k, v) in enumerate(codes.most_common()): put(ws, r0 + 1 + i, [k, " · ".join(CODE_DIM.get(c, c) for c in k), v, pct(v, len(ends))], [None, None, "#,##0", "0.0%"])
    widths(ws, [12, 30, 10, 10])

    # 军师
    ws = sheet("军师", "大家最常听谁的（玩到结局的局）")
    advs = Counter(x.get("adv", "") for x in ends if x.get("adv"))
    head(ws, 4, ["角色", "当军师的局数", "占比", "这些局的平均年化超额"])
    for i, (k, v) in enumerate(advs.most_common()):
        ann = [x["ann"] for x in ends if x.get("adv") == k and not x.get("dead")]
        put(ws, 5 + i, [META["ch"].get(k, k), v, pct(v, len(ends)), st.mean(ann) if ann else None], [None, "#,##0", "0.0%", "0.00%"])
    widths(ws, [12, 14, 10, 22])

    # 成绩
    ws = sheet("成绩分布", "十年成绩分布（玩到结局的局）")
    bins = [(-1, -.05), (-.05, -.03), (-.03, -.01), (-.01, 0), (0, .01), (.01, .03), (.03, .05), (.05, .07), (.07, 1)]
    head(ws, 4, ["年化超额区间", "局数", "占比"])
    put(ws, 5, ["中途清盘", len(dead), pct(len(dead), len(ends))], [None, "#,##0", "0.0%"])
    for i, (lo, hi) in enumerate(bins):
        c = sum(1 for x in full if lo <= x["ann"] < hi)
        lab = f"{'<' if lo == -1 else ''}{lo*100:.0f}% ~ {hi*100:.0f}%" if lo > -1 and hi < 1 else (f"< {hi*100:.0f}%" if lo == -1 else f"≥ {lo*100:.0f}%")
        put(ws, 6 + i, [lab, c, pct(c, len(ends))], [None, "#,##0", "0.0%"])
    dy = Counter(x.get("y") for x in dead)
    r0 = 8 + len(bins); head(ws, r0, ["清盘发生在", "局数"])
    for i, (k, v) in enumerate(sorted(dy.items(), key=lambda t: t[0] or 0)): put(ws, r0 + 1 + i, [f"{GY0 + (k or 1) - 1} 年（{META['reals'][(k or 1) - 1]}）", v])
    widths(ws, [30, 10, 10])

    # 思考时长
    ws = sheet("思考时长", "每一年的决定，大家想了多久", "从读完新闻到点击「执行」的秒数；超过 10 分钟的视为挂机，不计入")
    head(ws, 4, ["游戏年份", "真实事件", "中位数（秒）", "平均（秒）", "样本"])
    for y in range(10):
        v = [x["yt"][y] for x in rows if len(x.get("yt") or []) > y and x["yt"][y] is not None and x["yt"][y] < 600]
        put(ws, 5 + y, [f"{GY0 + y}", META["reals"][y], st.median(v) if v else None, st.mean(v) if v else None, len(v)], [None, None, "0", "0.0", "#,##0"])
    widths(ws, [10, 22, 14, 12, 8])

    # 下期素材
    ws = sheet("下期素材", "可以直接拿来讲的数字", "自动生成，发布前请核对口径；不要据此给出任何买卖建议")
    lines = []
    for y, c, t in yearpick:
        if not t: continue
        top, k = c.most_common(1)[0]; b = "ABC"[META["best"][y]]
        if top != b and k / t >= .4:
            lines.append(f"{GY0 + y} 年（真实是 {META['reals'][y]}）：最多人选「{META['opts'][y]['ABC'.index(top)]}」（{k/t:.0%}），但事后看最优的是「{META['opts'][y]['ABC'.index(b)]}」，只有 {pct(c.get(b,0), t):.0%} 的人选对。")
    if base: k, v = base.most_common(1)[0]; lines.append(f"最常见的投资人格是「{META['persona'].get(k, k)}」，占 {pct(v, len(ends)):.0%}。")
    if advs: k, v = advs.most_common(1)[0]; lines.append(f"最多人把 {META['ch'].get(k, k)} 当军师（{pct(v, len(ends)):.0%}）。")
    if ends:
        lines.append(f"{pct(len(dead), len(ends)):.0%} 的玩家没撑满十年，产品中途清盘。")
        lines.append(f"能守住张阿姨 10 万本金的玩家只有 {pct(sum(1 for r in ends if (r.get('aunt') or 0) >= 10), len(ends)):.0%}。")
    c1 = sum(v for k, v in reach.items() if k < 3)
    lines.append(f"{pct(c1, n_all):.0%} 的人在前三年就离开了游戏。")
    for i, l in enumerate(lines): ws.cell(row=4 + i, column=1, value=f"{i+1}. {l}").font = N
    widths(ws, [120])

    # 原始数据
    ws = sheet("原始数据", "每局一行")
    cols = ["day", "id", "ev", "src", "dev", "run", "ch", "y", "dead", "ann", "aunt", "rank", "beat", "code", "adv", "dur", "pro", "yt", "v"]
    head(ws, 3, cols)
    for i, x in enumerate(sorted(rows, key=lambda t: t.get("ts", 0))):
        put(ws, 4 + i, [json.dumps(x.get(c), ensure_ascii=False) if isinstance(x.get(c), list) else x.get(c) for c in cols])
    fn = f"玩家数据报告_{datetime.now():%Y%m%d}.xlsx"; wb.save(fn); print("报告已生成：", fn)

if __name__ == "__main__":
    main()
